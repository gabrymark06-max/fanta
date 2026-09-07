"""L'import di un listino esterno (FantaLab e simili).

Il formato esatto del file non e' noto in anticipo e cambia da tool a tool:
per questo il parser riconosce le colonne dai nomi invece di contare le
posizioni, e per questo questi test provano fogli scritti in modi diversi.
Quello che non deve mai succedere e' un import che sbaglia in silenzio.
"""

from __future__ import annotations

import csv
import io

import pytest

from fantabot.dati.esterno import importa, prezzi_esterni, riassunto


def _csv(righe: list[list[object]]) -> bytes:
    buffer = io.StringIO()
    csv.writer(buffer, delimiter=";").writerows(righe)
    return buffer.getvalue().encode("utf-8")


def _xlsx(righe: list[list[object]]) -> bytes:
    from openpyxl import Workbook

    libro = Workbook()
    for riga in righe:
        libro.active.append(riga)
    buffer = io.BytesIO()
    libro.save(buffer)
    return buffer.getvalue()


@pytest.fixture
def conn_listone(conn):
    conn.executemany(
        """
        INSERT INTO giocatori (id_fc, nome, nome_cerca, squadra, ruolo,
                               quota_iniziale, quota_attuale, fvm)
        VALUES (?, ?, ?, ?, ?, ?, ?, ?)
        """,
        [
            (1, "Malen", "malen", "ROM", "a", 38, 38, 450),
            (2, "Martinez L.", "martinez l", "INT", "a", 33, 33, 361),
            (3, "Dimarco", "dimarco", "INT", "d", 31, 31, 240),
            (4, "Svilar", "svilar", "ROM", "p", 19, 19, 85),
            (5, "Vlahovic", "vlahovic", "JUV", "a", 20, 20, 120),
        ],
    )
    conn.commit()
    return conn


def test_importa_un_foglio_con_prezzi_medi(conn_listone):
    dati = _csv(
        [
            ["Nome", "Squadra", "Ruolo", "Fascia", "Prezzo medio asta", "Prezzo max"],
            ["Malen", "ROM", "A", 1, 118, 140],
            ["Martinez L.", "INT", "A", 1, 96, 120],
            ["Dimarco", "INT", "D", 1, 72, 85],
        ]
    )
    esito = importa(conn_listone, dati, "strategia.csv")
    assert esito.riuscito
    assert esito.abbinati == 3
    assert esito.con_prezzo == 3
    assert esito.con_fascia == 3
    assert prezzi_esterni(conn_listone)[1] == 118


def test_riconosce_le_colonne_ovunque_siano(conn_listone):
    """Nessuna posizione fissa: conta cosa c'e' scritto nell'intestazione."""
    dati = _csv(
        [
            ["Fascia", "Costo medio", "Calciatore", "Team"],
            [2, 45, "Vlahovic", "JUV"],
            [1, 118, "Malen", "ROM"],
        ]
    )
    esito = importa(conn_listone, dati, "altro.csv")
    assert esito.riuscito
    assert esito.colonne_usate["nome"] == "Calciatore"
    assert esito.colonne_usate["prezzo_medio"] == "Costo medio"
    assert prezzi_esterni(conn_listone)[5] == 45


def test_salta_le_righe_di_intestazione_sopra_la_tabella(conn_listone):
    """I fogli esportati hanno spesso un titolo o una riga vuota in cima."""
    dati = _csv(
        [
            ["La mia strategia asta 2026/27", "", ""],
            ["", "", ""],
            ["Nome", "Squadra", "Prezzo medio"],
            ["Svilar", "ROM", 28],
        ]
    )
    esito = importa(conn_listone, dati, "export.csv")
    assert esito.riuscito
    assert prezzi_esterni(conn_listone)[4] == 28


def test_abbina_i_nomi_scritti_in_un_altro_modo(conn_listone):
    """Ogni tool scrive i nomi a modo suo, la squadra fa da giudice."""
    dati = _csv(
        [
            ["Nome", "Squadra", "Prezzo medio"],
            ["Lautaro Martinez", "INT", 96],
        ]
    )
    esito = importa(conn_listone, dati, "x.csv")
    assert esito.abbinati == 1
    assert prezzi_esterni(conn_listone)[2] == 96


def test_dice_quali_nomi_non_ha_riconosciuto(conn_listone):
    dati = _csv(
        [
            ["Nome", "Squadra", "Prezzo medio"],
            ["Malen", "ROM", 118],
            ["Chi Sono Io", "XXX", 40],
        ]
    )
    esito = importa(conn_listone, dati, "x.csv")
    assert esito.abbinati == 1
    assert "Chi Sono Io" in esito.non_trovati


def test_legge_anche_excel(conn_listone):
    dati = _xlsx(
        [
            ["Nome", "Squadra", "Prezzo medio"],
            ["Malen", "ROM", 118],
            ["Dimarco", "INT", 72],
        ]
    )
    esito = importa(conn_listone, dati, "strategia.xlsx")
    assert esito.riuscito
    assert esito.abbinati == 2


def test_un_foglio_senza_nomi_viene_rifiutato_spiegando(conn_listone):
    dati = _csv([["Colonna A", "Colonna B"], [1, 2]])
    esito = importa(conn_listone, dati, "x.csv")
    assert not esito.riuscito
    assert "nomi dei giocatori" in esito.errore
    # E dice cosa ha letto, cosi' si capisce cosa non e' andato.
    assert "Colonna A" in esito.errore


def test_un_foglio_coi_soli_nomi_non_serve_a_niente(conn_listone):
    dati = _csv([["Nome", "Squadra"], ["Malen", "ROM"]])
    esito = importa(conn_listone, dati, "x.csv")
    assert not esito.riuscito
    assert "prezzi o fasce" in esito.errore


def test_un_file_illeggibile_non_esplode(conn_listone):
    esito = importa(conn_listone, b"\x00\x01\x02non sono un foglio", "rotto.xlsx")
    assert not esito.riuscito
    assert esito.errore


def test_reimportare_aggiorna_senza_perdere_quello_che_c_era(conn_listone):
    importa(
        conn_listone,
        _csv([["Nome", "Prezzo medio"], ["Malen", 118]]),
        "a.csv",
    )
    importa(
        conn_listone,
        _csv([["Nome", "Fascia"], ["Malen", 1]]),
        "b.csv",
    )
    r = conn_listone.execute(
        "SELECT prezzo_medio, fascia FROM mercato_esterno WHERE id_fc = 1"
    ).fetchone()
    assert r["prezzo_medio"] == 118  # non cancellato dal secondo foglio
    assert r["fascia"] == 1


def test_il_riassunto_racconta_cosa_c_e(conn_listone):
    assert (riassunto(conn_listone).get("righe") or 0) == 0
    importa(
        conn_listone,
        _csv([["Nome", "Prezzo medio"], ["Malen", 118], ["Dimarco", 72]]),
        "a.csv",
    )
    r = riassunto(conn_listone)
    assert r["righe"] == 2
    assert r["prezzi"] == 2


def _xlsx_multi(fogli: dict[str, list[list[object]]]) -> bytes:
    from openpyxl import Workbook

    libro = Workbook()
    libro.remove(libro.active)
    for nome, righe in fogli.items():
        foglio = libro.create_sheet(nome)
        for riga in righe:
            foglio.append(riga)
    buffer = io.BytesIO()
    libro.save(buffer)
    return buffer.getvalue()


# Il formato vero di FantaLab: un foglio per ruolo, la colonna dei prezzi si
# chiama solo "Prezzo", e le fasce sono parole.
INTESTAZIONE_FANTALAB = [
    "Obiett.", "Fascia", "Ruolo", "Team", "Nome", "Prezzo", "PMA", "Quo",
    "Titolarità", "Affidabilità", "Integrità", "Commento", "Nota 1", "Nota 2",
]


def _riga_fantalab(nome, team, ruolo, fascia, prezzo, pma="9.2%", obiettivo=""):
    return [obiettivo, fascia, ruolo, team, nome, prezzo, pma, 19,
            5, 5, 5, "", "titolarissimo", ""]


def test_legge_il_formato_vero_di_fantalab(conn_listone):
    """Quattro fogli, colonna "Prezzo", fasce a parole: il formato reale.

    La prima versione dell'importatore avrebbe letto solo il primo foglio e
    non avrebbe riconosciuto nessuna colonna dei prezzi.
    """
    dati = _xlsx_multi(
        {
            "P": [INTESTAZIONE_FANTALAB, _riga_fantalab("Svilar", "ROM", "P", "Top", 55)],
            "D": [INTESTAZIONE_FANTALAB, _riga_fantalab("Dimarco", "INT", "D", "Top", 75)],
            "C": [INTESTAZIONE_FANTALAB],
            "A": [
                INTESTAZIONE_FANTALAB,
                _riga_fantalab("Malen", "ROM", "A", "Top", 180, pma="48.1%"),
                _riga_fantalab("Martinez L.", "INT", "A", "Semi-Top", 150, pma="40%"),
                _riga_fantalab("Vlahovic", "JUV", "A", "Terza", 40, obiettivo="X"),
            ],
        }
    )
    esito = importa(conn_listone, dati, "Strategia CarmySpecial.xlsx")
    assert esito.riuscito
    # Tutti e cinque, presi da fogli diversi.
    assert esito.abbinati == 5
    # "PMA" e' quanto viene pagato davvero, "Prezzo" e' il consiglio del tool:
    # due cose diverse, e il motore deve usare la prima.
    assert esito.colonne_usate["prezzo_medio"] == "PMA"
    assert esito.colonne_usate["prezzo_consigliato"] == "Prezzo"
    r = conn_listone.execute(
        "SELECT prezzo_medio, prezzo_consigliato FROM mercato_esterno WHERE id_fc = 1"
    ).fetchone()
    assert r["prezzo_medio"] == 48.1
    assert r["prezzo_consigliato"] == 180


def test_le_intestazioni_ripetute_non_finiscono_fra_gli_scarti(conn_listone):
    """Ogni foglio ripete la sua intestazione: non sono nomi sconosciuti."""
    dati = _xlsx_multi(
        {
            "P": [INTESTAZIONE_FANTALAB, _riga_fantalab("Svilar", "ROM", "P", "Top", 55)],
            "A": [INTESTAZIONE_FANTALAB, _riga_fantalab("Malen", "ROM", "A", "Top", 180)],
        }
    )
    esito = importa(conn_listone, dati, "x.xlsx")
    assert esito.non_trovati == []
    assert esito.righe_lette == 2


def test_le_fasce_scritte_a_parole_diventano_numeri(conn_listone):
    dati = _xlsx_multi(
        {
            "A": [
                INTESTAZIONE_FANTALAB,
                _riga_fantalab("Malen", "ROM", "A", "Top", 180),
                _riga_fantalab("Martinez L.", "INT", "A", "Semi-Top", 150),
                _riga_fantalab("Vlahovic", "JUV", "A", "Outsider", 12),
                _riga_fantalab("Dimarco", "INT", "D", "Non Impostata", 75),
            ]
        }
    )
    importa(conn_listone, dati, "x.xlsx")
    fasce = {
        r["id_fc"]: r["fascia"]
        for r in conn_listone.execute("SELECT id_fc, fascia FROM mercato_esterno")
    }
    assert fasce[1] == 1  # Top
    assert fasce[2] == 2  # Semi-Top
    assert fasce[5] == 5  # Outsider
    assert fasce[3] is None  # Non impostata


def test_una_colonna_prezzo_max_non_viene_scambiata_per_il_prezzo(conn_listone):
    dati = _csv(
        [
            ["Nome", "Prezzo max", "Prezzo medio"],
            ["Malen", 200, 118],
        ]
    )
    esito = importa(conn_listone, dati, "x.csv")
    assert esito.colonne_usate["prezzo_medio"] == "Prezzo medio"
    assert esito.colonne_usate["prezzo_max"] == "Prezzo max"
    assert prezzi_esterni(conn_listone)[1] == 118


def test_il_prezzo_osservato_batte_quello_consigliato(conn_listone):
    """PMA e Prezzo dicono due cose diverse, e il motore usa quella osservata.

    Il prezzo consigliato e' un'opinione del tool; il PMA e' quanto quel
    giocatore e' costato davvero, misurato sulle aste vere. Quando ci sono
    entrambi non c'e' partita.
    """
    dati = _csv(
        [
            ["Nome", "Prezzo", "PMA"],
            ["Malen", 180, "48.1%"],
        ]
    )
    importa(conn_listone, dati, "x.csv")
    assert prezzi_esterni(conn_listone)[1] == 48.1


def test_col_solo_prezzo_consigliato_si_usa_quello(conn_listone):
    """Non tutti i tool danno il prezzo osservato: meglio un'opinione di niente."""
    dati = _csv([["Nome", "Prezzo"], ["Malen", 180]])
    importa(conn_listone, dati, "x.csv")
    assert prezzi_esterni(conn_listone)[1] == 180


# -- i tre giudizi da 1 a 5 ----------------------------------------------


def test_legge_titolarita_integrita_e_affidabilita(conn_listone):
    """Le colonne dell'export FantaLab, con le loro intestazioni vere."""
    dati = _csv(
        [
            ["Nome", "Squadra", "Prezzo", "Titolarità", "Affidabilità", "Integrità"],
            ["Vlahovic", "JUV", 60, 5, 4, 1],
        ]
    )
    esito = importa(conn_listone, dati, "strategia.csv")
    assert esito.con_giudizi == 1
    r = conn_listone.execute(
        "SELECT titolarita, integrita, affidabilita FROM mercato_esterno"
    ).fetchone()
    assert (r["titolarita"], r["integrita"], r["affidabilita"]) == (5, 1, 4)


def test_un_numero_fuori_scala_non_e_un_giudizio(conn_listone):
    """Una colonna che contiene 87 e' un'altra colonna, non un refuso."""
    dati = _csv(
        [["Nome", "Squadra", "Prezzo", "Titolarità"], ["Vlahovic", "JUV", 60, 87]]
    )
    importa(conn_listone, dati, "strategia.csv")
    r = conn_listone.execute("SELECT titolarita FROM mercato_esterno").fetchone()
    assert r["titolarita"] is None
